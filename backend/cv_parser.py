import io
import pymupdf as fitz
import docx  # python-docx

class CVParseError(Exception):
    pass

def extract_text_from_pdf(content: bytes) -> str:
    """
    Extract readable text from PDF bytes using PyMuPDF.
    """
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        text_chunks = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text:
                text_chunks.append(page_text)
        doc.close()
        full_text = "\n".join(text_chunks).strip()
        return full_text
    except Exception as e:
        raise CVParseError(f"Gagal membaca format PDF: {str(e)}")

def extract_text_from_docx(content: bytes) -> str:
    """
    Extract readable text from DOCX bytes using python-docx.
    """
    try:
        file_stream = io.BytesIO(content)
        doc = docx.Document(file_stream)
        text_chunks = []
        
        # Read paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_chunks.append(para.text.strip())
                
        # Read tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_chunks.append(" | ".join(row_texts))
                    
        full_text = "\n".join(text_chunks).strip()
        return full_text
    except Exception as e:
        raise CVParseError(f"Gagal membaca format DOCX: {str(e)}")

def parse_cv_file(filename: str, content: bytes) -> str:
    """
    Validates and extracts real raw text from uploaded CV file.
    """
    MAX_SIZE = 5 * 1024 * 1024  # 5 MB (FR-01)
    if len(content) > MAX_SIZE:
        raise CVParseError("Your CV is too large. Maximum file size is 5 MB.")

    if not content or len(content) == 0:
        raise CVParseError("The uploaded file is empty. Please upload another file.")

    fname = filename.lower()
    if fname.endswith(".pdf"):
        raw_text = extract_text_from_pdf(content)
    elif fname.endswith(".docx"):
        raw_text = extract_text_from_docx(content)
    else:
        raise CVParseError("Unsupported file format. Please upload PDF or DOCX.")

    if not raw_text or len(raw_text.strip()) < 10:
        raise CVParseError("We couldn't extract text from this CV. Please upload another file.")

    return raw_text
